"""Give every video in a PowerPoint deck its own "click to play" cover.

Step 1 - list the videos and write a covers list to check:

    python3 set_covers.py "Tutorial deck.pptx"

  This writes covers.csv next to the deck: one row per video with the
  label and title its cover will show, guessed from the video's file name
  and the slide title. Fix any row that reads wrong.

Step 2 - make the new deck:

    python3 set_covers.py "Tutorial deck.pptx" covers.csv

  This saves "Tutorial deck - with covers.pptx" beside the original (the
  original is not changed) and a PNG of each cover in a covers folder.
"""

import csv
import io
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

from make_cover import render_cover

FIELDS = ["slide", "video", "file", "label", "title"]
GROUPS = {"provider": "Providers", "staff": "Staff", "member": "Members", "admin": "Admins"}


def videos_in(prs):
    """Yield (slide number, video number on that slide, slide, <p:pic>) for every video."""
    for n, slide in enumerate(prs.slides, start=1):
        pics = slide.shapes._spTree.xpath(".//p:pic[p:nvPicPr/p:nvPr/a:videoFile]")
        for i, pic in enumerate(pics, start=1):
            yield n, i, slide, pic


def video_file_name(slide, pic):
    # PowerPoint names a video after the file it was inserted from, while the
    # copy inside the deck is renamed media1.mp4, media2.mp4 and so on.
    name = pic.find(qn("p:nvPicPr")).find(qn("p:cNvPr")).get("name", "")
    if name and not re.fullmatch(r"(Movie|Video|Picture|Media)\s*\d*", name, re.I):
        return name
    rid = pic.xpath("./p:nvPicPr/p:nvPr/a:videoFile")[0].get(qn("r:link"))
    rel = slide.part.rels.get(rid)
    if rel is None:
        return ""
    target = rel.target_ref if rel.is_external else rel.target_part.partname
    return Path(str(target)).name


def slide_title(slide):
    title = slide.shapes.title
    return title.text_frame.text.strip() if title is not None and title.has_text_frame else ""


def guess(file_name, title, number):
    # "Provider 03 - Register your program.mp4" -> "Providers · Video 03", "Register your program"
    stem = re.sub(r"\.(mp4|m4v|mov|wmv|avi)$", "", file_name, flags=re.I).replace("_", " ")
    label = f"Video {number:02d}"
    m = re.search(r"(provider|staff|member|admin)s?\D{0,3}(\d+)", stem, re.I)
    if m:
        label = f"{GROUPS[m.group(1).lower()]} · Video {int(m.group(2)):02d}"
        rest = stem[m.end():].strip(" -–:").replace("-", " ")
        if rest and not title:
            title = rest[:1].upper() + rest[1:]
    return label, title or stem


def write_list(deck, csv_path):
    prs = Presentation(deck)
    rows = []
    for k, (n, i, slide, pic) in enumerate(videos_in(prs), start=1):
        name = video_file_name(slide, pic)
        label, title = guess(name, slide_title(slide), k)
        rows.append({"slide": n, "video": i, "file": name, "label": label, "title": title})
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    print(f"{len(rows)} video(s) found. Check the labels and titles in {csv_path}:")
    for r in rows:
        print(f"  slide {r['slide']:>3}  {r['label']:<24} {r['title']}   ({r['file']})")


def apply(deck, csv_path):
    with open(csv_path, newline="", encoding="utf-8-sig") as f:
        wanted = {(int(r["slide"]), int(r["video"])): r for r in csv.DictReader(f)}
    prs = Presentation(deck)
    out_dir = deck.with_name("covers")
    out_dir.mkdir(exist_ok=True)
    done = 0
    for n, i, slide, pic in videos_in(prs):
        row = wanted.get((n, i))
        if row is None:
            print(f"  slide {n}: video {i} is not in {csv_path.name}; left as is")
            continue
        # Draw the cover at the video's own shape so it is not stretched.
        ext = pic.find(qn("p:spPr")).find(qn("a:xfrm")).find(qn("a:ext"))
        ratio = int(ext.get("cx")) / int(ext.get("cy"))
        width, height = (1920, round(1920 / ratio)) if ratio >= 16 / 9 else (round(1080 * ratio), 1080)
        image = render_cover(row["label"], row["title"], width, height)
        png = io.BytesIO()
        image.save(png, "PNG", optimize=True)
        safe = re.sub(r"[^\w -]", "", f"{n:02d} {row['title']}").strip()
        image.save(out_dir / f"{safe}.png", optimize=True)

        blip = pic.find(qn("p:blipFill")).find(qn("a:blip"))
        old = blip.get(qn("r:embed"))
        _, new = slide.part.get_or_add_image_part(io.BytesIO(png.getvalue()))
        blip.set(qn("r:embed"), new)
        if old and old != new:
            slide.part.drop_rel(old)
        done += 1
        print(f"  slide {n:>3}  {row['label']:<24} {row['title']}")
    out = deck.with_name(f"{deck.stem} - with covers.pptx")
    prs.save(out)
    print(f"{done} cover(s) set. Saved {out}")


if __name__ == "__main__":
    if len(sys.argv) == 2:
        deck = Path(sys.argv[1])
        write_list(deck, deck.with_name("covers.csv"))
    elif len(sys.argv) == 3:
        apply(Path(sys.argv[1]), Path(sys.argv[2]))
    else:
        sys.exit(__doc__)
